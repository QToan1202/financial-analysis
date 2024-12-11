import io
from io import BytesIO
from uuid import uuid4
from langchain.text_splitter import RecursiveCharacterTextSplitter
from langchain_core.documents import Document
from typing import List
from fastapi import HTTPException


class FileService:
    def __init__(self, content: bytes, file_extension=str) -> None:
        self.content = content
        self.file_extension = file_extension
        self._id_for_doc = str(uuid4())

    def _split_documents(self, page_content, *, metadata: dict) -> List[Document]:
        doc = [Document(page_content=page_content,
                        metadata=metadata | {"id": self._id_for_doc})]
        text_splitter = RecursiveCharacterTextSplitter.from_tiktoken_encoder(
            chunk_size=400, chunk_overlap=10, separators=["\n\n", "\n", ".", " "], keep_separator=False,)
        splits = text_splitter.split_documents(doc)

        return splits

    def _handle_pdf(self, *, metadata: dict) -> List[Document]:
        from pypdf import PdfReader

        pdf_content = io.BytesIO(self.content)
        pdf_reader = PdfReader(pdf_content)
        pdf_text = ""
        for page in pdf_reader.pages:
            pdf_text += page.extract_text()

        return self._split_documents(page_content=pdf_text, metadata=metadata)

    def _handle_txt(self, *, metadata: dict) -> List[Document]:
        file_text = self.content.decode("utf-8")

        return self._split_documents(page_content=file_text, metadata=metadata)

    def _handle_docx(self, *, metadata: dict) -> List[Document]:
        from docx import Document as DocxDocument

        file_stream = BytesIO(self.content)
        document = DocxDocument(file_stream)
        file_text = "\n".join(
            paragraph.text for paragraph in document.paragraphs)

        return self._split_documents(page_content=file_text, metadata=metadata)

    def handle_split_file_content(self, *, metadata: dict) -> List[Document]:
        if (self.file_extension == "txt"):
            return self._handle_txt(metadata=metadata)
        elif (self.file_extension == "pdf"):
            return self._handle_pdf(metadata=metadata)
        elif (self.file_extension == "docx"):
            return self._handle_docx(metadata=metadata)
        else:
            raise HTTPException(
                status_code=415,
                detail=f"{self.file_extension.upper()} format is not directly supported."
            )
